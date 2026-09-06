"""Extract PowerPoint slide text and read it aloud."""

from __future__ import annotations

from pathlib import Path

import cv2
import numpy as np
import pyttsx3
import pytesseract
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


class SlideReader:
    """Read slide text from PowerPoint files, with OCR as a fallback."""

    def __init__(self) -> None:
        self.engine = pyttsx3.init()

    @staticmethod
    def extract_text_from_slide(slide) -> str:
        """Extract slide text and OCR picture shapes when needed."""
        text_parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_parts.append(shape.text.strip())
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                # Decode the embedded picture in memory for the OCR fallback.
                image = cv2.imdecode(
                    np.frombuffer(shape.image.blob, dtype=np.uint8),
                    cv2.IMREAD_COLOR,
                )
                if image is not None:
                    ocr_text = pytesseract.image_to_string(image).strip()
                    if ocr_text:
                        text_parts.append(ocr_text)
        return "\n".join(text_parts)

    @staticmethod
    def extract_text_from_image(image_path: str | Path) -> str:
        """Use OCR to read text contained in an image file."""
        image = cv2.imread(str(image_path))
        if image is None:
            raise FileNotFoundError(f"Could not read image: {image_path}")
        return pytesseract.image_to_string(image).strip()

    def read_presentation(self, presentation_path: str | Path) -> list[str]:
        """Extract and speak text from every slide in a presentation."""
        presentation = Presentation(str(presentation_path))
        slide_text = [self.extract_text_from_slide(slide) for slide in presentation.slides]
        for index, text in enumerate(slide_text, start=1):
            if text:
                print(f"Slide {index}: {text}")
                self.speak(text)
        return slide_text

    def speak(self, text: str) -> None:
        """Speak text through the system's configured audio output."""
        self.engine.say(text)
        self.engine.runAndWait()


if __name__ == "__main__":
    path = input("Enter a .pptx file path: ").strip()
    SlideReader().read_presentation(path)
